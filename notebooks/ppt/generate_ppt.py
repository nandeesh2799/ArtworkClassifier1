"""
PowerPoint Presentation Generator for Artwork Classification Project.

Usage:
  python3 notebooks/ppt/generate_ppt.py

Output:
  artwork_classification_presentation.pptx

Tip:
  After training, figures will be saved to `model/exports/`:
  - training_curves.png
  - confusion_matrix.png
  You can manually insert them into the PPT or extend this script to embed them.
"""

# (Moved from project root to keep structure clean.)

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from datetime import datetime


class Colors:
    PRIMARY = RGBColor(46, 134, 171)  # Blue
    SECONDARY = RGBColor(162, 59, 114)  # Purple
    DARK = RGBColor(44, 62, 80)
    TEXT = RGBColor(51, 51, 51)
    TEXT_LIGHT = RGBColor(108, 117, 125)


class ArtworkPPTGenerator:
    def __init__(self):
        self.prs = Presentation()
        self.prs.slide_width = Inches(13.333)  # 16:9
        self.prs.slide_height = Inches(7.5)
        self.slide_count = 0

    def _add_title(self, slide, title_text: str) -> None:
        title = slide.shapes.title
        title.text = title_text
        p = title.text_frame.paragraphs[0]
        p.font.size = Pt(36)
        p.font.bold = True
        p.font.color.rgb = Colors.PRIMARY

    def _add_body(self, slide, text: str, font_size: int = 18) -> None:
        body = slide.placeholders[1]
        body.text = text
        for par in body.text_frame.paragraphs:
            par.font.size = Pt(font_size)
            par.font.color.rgb = Colors.TEXT
            if par.text.strip().endswith(":"):
                par.font.bold = True
                par.font.color.rgb = Colors.PRIMARY

    def add_title_slide(self):
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[0])
        self.slide_count += 1

        title = slide.shapes.title
        title.text = "Artwork Classification Using Deep Learning"
        p = title.text_frame.paragraphs[0]
        p.font.size = Pt(48)
        p.font.bold = True
        p.font.color.rgb = Colors.PRIMARY

        subtitle = slide.placeholders[1]
        subtitle.text = "Transfer Learning (EfficientNet / ResNet)\nStreamlit Web App + Full Evaluation"
        subtitle.text_frame.paragraphs[0].font.size = Pt(26)
        subtitle.text_frame.paragraphs[0].font.color.rgb = Colors.SECONDARY

        tx = slide.shapes.add_textbox(Inches(1), Inches(5.6), Inches(11), Inches(1.2)).text_frame
        tx.word_wrap = True
        p1 = tx.add_paragraph()
        p1.text = "Presented by: [Your Name]"
        p1.font.size = Pt(20)
        p1.font.bold = True
        p1.font.color.rgb = Colors.DARK
        p1.alignment = PP_ALIGN.CENTER

        p2 = tx.add_paragraph()
        p2.text = f"Date: {datetime.now().strftime('%B %d, %Y')}"
        p2.font.size = Pt(14)
        p2.font.color.rgb = Colors.TEXT_LIGHT
        p2.alignment = PP_ALIGN.CENTER

    def add_problem_slide(self):
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[1])
        self.slide_count += 1
        self._add_title(slide, "Problem & Dataset")
        self._add_body(
            slide,
            """Problem:
• Automatic artwork type classification (drawings / paintings / sculptures / engravings)
• Useful for museums, digital archives, education, and search

Primary Dataset (Mandatory):
• Kaggle: thedownhill/art-images-drawings-painting-sculpture-engraving
• Cleaned corrupted images before training
• Stratified train/val/test split""",
        )

    def add_model_slide(self):
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[1])
        self.slide_count += 1
        self._add_title(slide, "Model")
        self._add_body(
            slide,
            """Transfer Learning:
• Backbone: EfficientNetB0 (default) or ResNet50
• Frozen base → train head → fine-tune top layers

Head:
• GlobalAveragePooling → Dense → Dropout → Softmax""",
        )

    def add_results_slide(self):
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[1])
        self.slide_count += 1
        self._add_title(slide, "Evaluation")
        self._add_body(
            slide,
            """Metrics:
• Accuracy
• Precision / Recall / F1 (weighted)
• Confusion Matrix

Artifacts saved to:
• model/exports/training_curves.png
• model/exports/confusion_matrix.png
• model/exports/test_metrics.json""",
        )

    def add_app_slide(self):
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[1])
        self.slide_count += 1
        self._add_title(slide, "Web Application (Streamlit)")
        self._add_body(
            slide,
            """Features:
• Upload image → prediction + confidence
• Clean UI with custom CSS
• Custom footer: “Final Year Project”
• Streamlit chrome minimized/hidden where possible""",
        )

    def add_qa_slide(self):
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[1])
        self.slide_count += 1
        self._add_title(slide, "Q&A")
        self._add_body(
            slide,
            """Thank you!

Questions?""",
            font_size=24,
        )

    def save(self, filename: str) -> str:
        self.prs.save(filename)
        return filename

    def generate(self) -> str:
        self.add_title_slide()
        self.add_problem_slide()
        self.add_model_slide()
        self.add_results_slide()
        self.add_app_slide()
        self.add_qa_slide()
        return self.save("artwork_classification_presentation.pptx")


def main() -> None:
    gen = ArtworkPPTGenerator()
    out = gen.generate()
    print("Saved:", out)


if __name__ == "__main__":
    main()

